# Copyright (c) 2023-present Plane Software, Inc. and contributors
# SPDX-License-Identifier: AGPL-3.0-only
# See the LICENSE file for details.

"""Roster parsing for the strict eleven-column member import template."""

import io

import pytest
from openpyxl import Workbook

from plane.research.services.accounts import AccountError
from plane.research.services.user_import import parse_advisors, parse_students, sheet_rows

pytestmark = pytest.mark.unit

ROSTER_HEADER = "姓名,学号,邮件,手机号,年级,人员类别,业务方向,小组,主导师,联合导师1,联合导师2\n"
ROSTER_ROW = (
    "邱智鑫,20220230156625,qiuzhixin@stu.xmu.edu.cn,17706010502,23,学生,基础研究,"
    "石墨负极小组,刘俊扬,陈志昕,白杰\n"
)


def test_new_template_headers_and_values_are_parsed():
    rows = parse_students((ROSTER_HEADER + ROSTER_ROW).encode("utf-8"), "roster.csv")

    assert len(rows) == 1
    row = rows[0]
    assert row.row_number == 2
    assert row.name == "邱智鑫"
    assert row.student_no == "20220230156625"
    assert row.email == "qiuzhixin@stu.xmu.edu.cn"
    assert row.phone == "17706010502"
    assert row.grade == "23"
    assert row.category == "STUDENT"
    assert row.business_category == "BASIC_RESEARCH"
    assert row.group == "石墨负极小组"
    assert row.primary_advisor_name == "刘俊扬"
    assert row.co_advisor_names == ("陈志昕", "白杰")


def test_new_template_allows_any_column_order_and_defaults_category():
    payload = (
        "主导师,小组,邮件,姓名,联合导师2,手机号,业务方向,学号,联合导师1,人员类别,年级\n"
        "刘俊扬,器件小组,alice@example.com,Alice,,,基础研究,S001,,,2026\n"
    )

    row = parse_students(payload.encode("utf-8"), "roster.csv")[0]

    assert row.name == "Alice"
    assert row.category == "STUDENT"
    assert row.co_advisor_names == ()


def test_gb18030_csv_is_decoded():
    rows = parse_students((ROSTER_HEADER + ROSTER_ROW).encode("gb18030"), "roster.csv")
    assert rows[0].name == "邱智鑫"


def test_xlsx_roster_is_parsed():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["姓名", "学号", "邮件", "手机号", "年级", "人员类别", "业务方向", "小组", "主导师", "联合导师1", "联合导师2"])
    sheet.append(["邱智鑫", "S001", "q@example.com", "177", "23", "学生", "基础研究", "器件小组", "刘俊扬", "", ""])
    buffer = io.BytesIO()
    workbook.save(buffer)

    rows = parse_students(buffer.getvalue(), "roster.xlsx")

    assert len(rows) == 1
    assert rows[0].email == "q@example.com"
    assert rows[0].phone == "177"


@pytest.mark.parametrize(("category", "degree"), [("Ph.D", "PHD"), ("MS", "MS")])
def test_supplied_template_aliases_are_parsed(category, degree):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["姓名", "学号", "邮件", "电话", "年级", "人员类别", "业务方向", "小组", "主导师", "联合导师1", "联合导师2", "备注"])
    sheet.append(["Alice", 20230156625, "alice@example.com", 17700000000, 23, category, "基础研究", "器件", "导师甲", "", "", ""])
    # Match the supplied workbook's styled but empty columns and rows.
    sheet.cell(216, 26).number_format = "@"
    buffer = io.BytesIO()
    workbook.save(buffer)

    rows = parse_students(buffer.getvalue(), "π-Lab学生-导入信息表.xlsx")

    assert len(rows) == 1
    assert rows[0].phone == "17700000000"
    assert rows[0].student_no == "20230156625"
    assert rows[0].grade == "23"
    assert rows[0].category == "STUDENT"
    assert rows[0].degree == degree
    assert rows[0].raw["category"] == category


def test_supplied_advisor_template_name_header_is_parsed():
    workbook = Workbook()
    workbook.active.append(["导师姓名", "邮箱"])
    workbook.active.append(["导师甲", " ADVISOR@example.com "])
    buffer = io.BytesIO()
    workbook.save(buffer)

    assert parse_advisors(buffer.getvalue(), "导师信息表.xlsx") == {
        "导师甲": "advisor@example.com"
    }


def test_duplicate_nonempty_headers_are_rejected():
    with pytest.raises(AccountError, match="重复"):
        parse_students((ROSTER_HEADER.rstrip() + ",姓名\n" + ROSTER_ROW).encode(), "roster.csv")


@pytest.mark.parametrize(
    "header",
    [
        "姓名,学号,邮件,手机号,年级,人员类别,业务方向,小组,主导师,联合导师1\n",
        "分组,姓名,学号,年级,学位,负责导师,邮件,电话\n",
        "name,student_no,email,phone,grade,category,business_category,group,primary_advisor,co_advisor_1,co_advisor_2\n",
    ],
)
def test_missing_new_header_or_legacy_templates_are_rejected(header):
    with pytest.raises(AccountError) as error:
        parse_students((header + "a,b,c,d,e,f,g,h,i,j,k\n").encode("utf-8"), "roster.csv")
    assert error.value.error_code == "user_import_file_invalid"


def test_empty_payload_is_rejected():
    with pytest.raises(AccountError) as error:
        sheet_rows(b"", "roster.csv")
    assert error.value.error_code == "user_import_file_required"


def test_advisor_table_is_normalised_and_exact_duplicates_are_allowed():
    payload = (
        "姓名,邮箱\n"
        "刘俊扬,liujunyang@xmu.edu.cn\n"
        " 刘俊扬 ,LIUJUNYANG@xmu.edu.cn\n"
        "陈志昕,chenzhixin@xmu.edu.cn\n"
    )

    advisors = parse_advisors(payload.encode("utf-8"), "advisors.csv")

    assert advisors == {
        "刘俊扬": "liujunyang@xmu.edu.cn",
        "陈志昕": "chenzhixin@xmu.edu.cn",
    }


def test_advisor_table_rejects_same_name_with_different_emails():
    payload = "姓名,邮箱\n刘俊扬,first@example.com\n 刘俊扬 ,second@example.com\n"

    with pytest.raises(AccountError) as error:
        parse_advisors(payload.encode("utf-8"), "advisors.csv")

    assert error.value.error_code == "user_import_file_invalid"
    assert "刘俊扬" in error.value.message


def test_advisor_table_rejects_duplicate_headers():
    payload = "导师姓名,邮箱,邮箱\n导师甲,first@example.com,second@example.com\n"

    with pytest.raises(AccountError, match="重复"):
        parse_advisors(payload.encode("utf-8"), "advisors.csv")


def test_duplicate_advisor_names_in_one_roster_row_are_rejected_by_row_validation(db):
    payload = ROSTER_HEADER + ROSTER_ROW.replace("刘俊扬,陈志昕,白杰", "刘俊扬,陈志昕,刘俊扬")
    row = parse_students(payload.encode("utf-8"), "roster.csv")[0]

    from plane.research.services.user_import import validate_row

    error_code, message = validate_row(row)
    assert error_code == "user_import_row_invalid"
    assert "导师不能重复" in message
